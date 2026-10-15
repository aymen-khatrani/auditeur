"""Lecteur generique d'extractions DOORS au format Excel.

Le template Alstom/IBM DOORS "Export to Excel" produit quatre onglets :
  Cover       : page de garde (projet, titre, version, baseline)
  MetaData    : contexte d'export (module path, baseline, vue, date)
  ModuleRich  : attributs du module (Name, Prefix, Description, ...)
  ViewRich    : une ligne par objet, colonnes = attributs de la vue exportee

Aucune colonne n'est codee en dur : le mapping vient de config.yaml, ce qui
permet de rejouer l'outil sur un autre projet / une autre vue DOORS.
"""
from __future__ import annotations

import os
import re
from typing import Dict, List, Optional, Sequence

import openpyxl

from .model import Corpus, Link, Module, Requirement

SEP_BLOCK = "┈"          # separateur d'entrees multiples dans les cellules DOORS
RE_LINK_HEADER = re.compile(r"^\s*(In|Out)\s*\[(.+?)\]\s*$", re.I)
RE_LINK_TARGET = re.compile(r"^\s*(?P<mod>.+?)#(?P<abs>\d+)\s*$")
RE_ID_ABS = re.compile(r"^(?P<prefix>.*?)[\s\-](?P<abs>\d+)$")
RE_HEADING_NUM = re.compile(r"^\s*(?P<num>\d+(?:\.\d+)*)\s+(?P<title>\S.*)$", re.S)
RE_REQ_ID = re.compile(r"[A-Za-z0-9][A-Za-z0-9_.\-]*[\s\-]\d+")


# --------------------------------------------------------------------------- utils
def _txt(value) -> str:
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").strip()


def _sheet_pairs(ws) -> Dict[str, str]:
    """Onglets cle/valeur (MetaData, ModuleRich)."""
    out: Dict[str, str] = {}
    for row in ws.iter_rows(values_only=True):
        if not row:
            continue
        key = _txt(row[0])
        if not key:
            continue
        val = _txt(row[1]) if len(row) > 1 else ""
        out[key] = val
    return out


def _pick(header: Sequence[str], names: Sequence[str]) -> Optional[int]:
    """Index de la premiere colonne du header correspondant a `names`."""
    lowered = [h.lower() for h in header]
    for name in names:
        n = name.lower()
        if n in lowered:
            return lowered.index(n)
    return None


def detect_level(levels: List[dict], *hints: str) -> str:
    """Determine le niveau de referentiel a partir du nom de module/fichier."""
    for hint in hints:
        if not hint:
            continue
        base = hint.strip()
        for lvl in levels:
            for pattern in lvl.get("detect", []):
                if re.search(pattern, base, re.I):
                    return lvl["id"]
    return "UNKNOWN"


def split_id(rid: str) -> tuple[str, Optional[int]]:
    """"SA.10.RS 1014" -> ("SA.10.RS", 1014) ; "PKP-ANNEX1.CUS-3339" -> (..., 3339)."""
    m = RE_ID_ABS.match(rid.strip())
    if not m:
        return rid.strip(), None
    return m.group("prefix"), int(m.group("abs"))


# --------------------------------------------------------------------------- liens
def parse_object_links(cell: str) -> List[Link]:
    """Decode la colonne ~ObjectLinks (blocs 'In [Specifies]' + cibles Module#123)."""
    links: List[Link] = []
    direction, ltype = "", ""
    for raw in cell.split("\n"):
        line = raw.strip()
        if not line:
            continue
        head = RE_LINK_HEADER.match(line)
        if head:
            direction, ltype = head.group(1).title(), head.group(2).strip()
            continue
        tgt = RE_LINK_TARGET.match(line)
        if tgt and direction:
            links.append(Link(direction=direction, link_type=ltype,
                              target_module=tgt.group("mod").strip(),
                              target_abs=int(tgt.group("abs"))))
    return links


def parse_requirement_refs(cell: str, source: str, direction: str) -> List[Link]:
    """Decode ParentRequirements / ChildRequirements.

    Format : "<ID>, <nom de module>\\n§section\\ntexte..." repete et separe par
    une ligne de caracteres '┈'.
    """
    links: List[Link] = []
    for block in re.split(f"[{SEP_BLOCK}]+", cell):
        block = block.strip()
        if not block:
            continue
        first = block.split("\n", 1)[0]
        if "," not in first:
            continue
        rid, module = first.split(",", 1)
        rid, module = rid.strip(), module.strip()
        if not RE_REQ_ID.fullmatch(rid or ""):
            continue
        _, abs_no = split_id(rid)
        links.append(Link(direction=direction, link_type=source,
                          target_module=module, target_abs=abs_no,
                          target_id=rid, source=source))
    return links


# --------------------------------------------------------------------------- parsing
def load_module(path: str, cfg: dict) -> Module:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        meta = _sheet_pairs(wb["MetaData"]) if "MetaData" in wb.sheetnames else {}
        rich = _sheet_pairs(wb["ModuleRich"]) if "ModuleRich" in wb.sheetnames else {}
        cover = _sheet_pairs(wb["Cover"]) if "Cover" in wb.sheetnames else {}

        fname = os.path.basename(path)
        name = rich.get("Name") or os.path.splitext(fname)[0]
        prefix = (rich.get("Prefix") or "").strip()
        level = detect_level(cfg["levels"], name, fname, meta.get("Module Path", ""))

        module = Module(
            name=name, prefix=prefix, level=level,
            path=meta.get("Module Path", ""),
            baseline=cover.get("DOORS Baseline") or meta.get("Baseline", ""),
            export_time=meta.get("UTC Export Time", ""),
            file=fname,
            meta={**{f"cover.{k}": v for k, v in cover.items()},
                  **{f"meta.{k}": v for k, v in meta.items()},
                  **{f"module.{k}": v for k, v in rich.items() if not k.startswith("~")}},
        )

        if "ViewRich" not in wb.sheetnames:
            return module

        ws = wb["ViewRich"]
        rows = ws.iter_rows(values_only=True)
        header = [_txt(c) for c in next(rows)]
        cols = {key: _pick(header, names) for key, names in cfg["columns"].items()}
        # colonnes 'parents' multiples : on les cumule toutes
        parent_cols = [i for i, h in enumerate(header)
                       if h.lower() in {n.lower() for n in cfg["columns"]["parents"]}]
        child_cols = [i for i, h in enumerate(header)
                      if h.lower() in {n.lower() for n in cfg["columns"]["children"]}]

        def cell(row, key) -> str:
            i = cols.get(key)
            if i is None or i >= len(row):
                return ""
            return _txt(row[i])

        section_stack: List[tuple[str, str]] = []   # (numero, titre)
        for rownum, row in enumerate(rows, start=2):
            rid = cell(row, "id")
            if not rid:
                continue
            text = cell(row, "text")
            obj_type = cell(row, "obj_type") or "<no object type>"
            prefix_r, abs_no = split_id(rid)

            if obj_type == "Heading":
                m = RE_HEADING_NUM.match(text)
                if m:
                    num, title = m.group("num"), m.group("title").split("\n")[0]
                    depth = num.count(".")
                    section_stack = section_stack[:depth]
                    section_stack.append((num, title.strip()))
            section = " > ".join(f"{n} {t}" for n, t in section_stack)

            links: List[Link] = parse_object_links(cell(row, "links"))
            for i in parent_cols:
                if i < len(row):
                    links += parse_requirement_refs(_txt(row[i]), "ParentRequirements", "In")
            for i in child_cols:
                if i < len(row):
                    links += parse_requirement_refs(_txt(row[i]), "ChildRequirements", "Out")

            module.requirements.append(Requirement(
                uid=f"{name}::{rid}", rid=rid, abs_no=abs_no, module=name, level=level,
                text=text, obj_type=obj_type, number=cell(row, "number"),
                section=section, allocated_to=cell(row, "allocated"),
                scope=cell(row, "scope"), rationale=cell(row, "rationale"),
                alloc_status=cell(row, "alloc_status"), demo=cell(row, "demo"),
                links=links, row=rownum,
            ))
        return module
    finally:
        wb.close()


def load_corpus(paths: List[str], cfg: dict) -> Corpus:
    corpus = Corpus()
    for p in sorted(paths):
        corpus.modules.append(load_module(p, cfg))
    _resolve_links(corpus)
    return corpus


def _resolve_links(corpus: Corpus) -> None:
    """Resout Module#abs -> ID DOORS pour les modules presents dans le corpus."""
    idx = corpus.module_index()
    abs_idx = {m.name: m.by_abs() for m in corpus.modules}
    for module in corpus.modules:
        for req in module.requirements:
            for link in req.links:
                if link.target_id:
                    continue
                tgt = idx.get(link.target_module) or idx.get(link.target_module.lower())
                if tgt is None or link.target_abs is None:
                    continue
                hit = abs_idx[tgt.name].get(link.target_abs)
                if hit is not None:
                    link.target_id = hit.rid


def discover(data_dir: str) -> List[str]:
    """Liste les extractions Excel d'un repertoire (ignore les fichiers temporaires)."""
    out = []
    for f in os.listdir(data_dir):
        if f.startswith("~$") or not f.lower().endswith((".xlsx", ".xlsm")):
            continue
        out.append(os.path.join(data_dir, f))
    return out
