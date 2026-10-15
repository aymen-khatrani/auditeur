"""Generation des livrables : matrice de tracabilite Excel + synthese Markdown."""
from __future__ import annotations

import datetime as dt
import json
import os
from typing import Dict, List, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import normalize as nz

STATUS_COLOR = {
    "LIEN_EXISTANT_CONFIRME": "C6EFCE",
    "LIEN_EXISTANT_A_REVOIR": "FFEB9C",
    "PROPOSITION": "DDEBF7",
    "A_VERIFIER": "FFF2CC",
    "ECART_SANS_CANDIDAT": "FFC7CE",
}
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(color="FFFFFF", bold=True)


def _sheet(wb: Workbook, title: str, header: Sequence[str], rows: Sequence[Sequence],
           widths: Sequence[int] = (), status_col: int = 0):
    ws = wb.create_sheet(title[:31])
    ws.append(list(header))
    for cell in ws[1]:
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append(list(row))
    if rows and status_col:
        for r in range(2, len(rows) + 2):
            color = STATUS_COLOR.get(str(ws.cell(r, status_col).value or ""))
            if color:
                ws.cell(r, status_col).fill = PatternFill("solid", fgColor=color)
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(rows)+1}"
    return ws


def _cut(text: str, n: int = 700) -> str:
    t = nz.clean(text)
    return t if len(t) <= n else t[: n - 1] + "…"


def write_excel(path: str, chains: List[Dict], cfg: dict, context: Dict) -> str:
    wb = Workbook()
    wb.remove(wb.active)

    # ------------------------------------------------------------- Synthese
    rows = []
    for ch in chains:
        s, rs, ev = ch["summary"], ch["reuse"], ch["eval"]
        rows.append([
            s["chaine"], s["exigences_amont"], s["deja_tracees"],
            s["taux_tracabilite_existant"], s["propositions_de_chainage"],
            s["taux_couverture_apres_proposition"], s["ecarts_sans_candidat"],
            s["couverture_partielle"], rs["taux_reuse_estime"], rs["quasi_identiques"],
            s["confiance_mediane_meilleur_candidat"],
            ev.get("rappel@5", ""), ev.get("precision@1", ""),
            ev.get("exigences_evaluables", 0),
        ])
    _sheet(wb, "Synthese",
           ["Chaine", "Exigences amont", "Deja tracees", "Taux tracabilite",
            "Propositions", "Couverture apres proposition", "Ecarts sans candidat",
            "Couverture partielle", "Taux de reuse estime", "Quasi-identiques",
            "Confiance mediane", "Rappel@5 (vs DOORS)", "Precision@1 (vs DOORS)",
            "Exigences evaluables"],
           rows, widths=[22, 14, 12, 12, 12, 16, 14, 14, 14, 13, 13, 14, 14, 14])

    # -------------------------------------------------------------- Matrice
    mrows = []
    for ch in chains:
        for res in ch["results"]:
            if not res.candidates:
                mrows.append([ch["label"], res.src.rid, res.src.module, res.src.section,
                              _cut(res.src.text), res.status, res.coverage,
                              "", "", "", "", 0.0, "non", "", "", "", "", "", "", "", ""])
                continue
            for rank, c in enumerate(res.candidates, 1):
                mrows.append([
                    ch["label"], res.src.rid, res.src.module, res.src.section,
                    _cut(res.src.text), res.status, res.coverage,
                    rank, c.tgt.rid, c.tgt.module, _cut(c.tgt.text),
                    round(c.confidence, 4), "oui" if c.existing else "non",
                    ", ".join(c.link_types), c.llm_verdict, c.llm_reason,
                    round(c.feats["emb"], 3), round(c.feats["tfidf"], 3),
                    round(c.feats["fuzzy"], 3), round(c.feats["numeric"], 3),
                    round(c.feats["section"], 3),
                ])
    _sheet(wb, "Matrice",
           ["Chaine", "ID amont", "Module amont", "Section amont", "Texte amont",
            "Statut", "Couverture", "Rang", "ID aval", "Module aval", "Texte aval",
            "Taux de confiance", "Lien existant", "Type de lien", "Verdict LLM",
            "Justification LLM", "sim embeddings", "sim tf-idf", "sim lexicale",
            "accord valeurs", "proximite section"],
           mrows, widths=[14, 20, 24, 30, 60, 24, 12, 6, 20, 24, 60, 13, 11, 18, 18,
                          50, 11, 11, 11, 11, 11], status_col=6)

    # --------------------------------------------------------------- Ecarts
    grows = []
    for ch in chains:
        for res in ch["results"]:
            if res.status == "LIEN_EXISTANT_CONFIRME" and res.coverage != "partielle":
                continue
            grows.append([
                ch["label"], res.src.rid, res.src.module, res.src.section,
                _cut(res.src.text), res.status, res.coverage,
                res.src.allocated_to, res.src.demo,
                " | ".join(_cut(a, 200) for a in res.missing_aspects),
                res.gap_note,
                res.candidates[0].tgt.rid if res.candidates else "",
                round(res.candidates[0].confidence, 4) if res.candidates else 0.0,
            ])
    _sheet(wb, "Ecarts",
           ["Chaine", "ID amont", "Module amont", "Section", "Texte amont", "Statut",
            "Couverture", "Alloue a", "Demonstration prevue", "Aspects non couverts",
            "Commentaire d'ecart", "Meilleur candidat", "Confiance"],
           grows, widths=[14, 20, 24, 28, 60, 24, 12, 26, 20, 60, 45, 20, 11],
           status_col=6)

    # ----------------------------------------------------------- Allocation
    arows = [[ch["label"]] + list(r.values()) for ch in chains for r in ch["allocation"]]
    if arows:
        _sheet(wb, "Allocation",
               ["Chaine", "ID amont", "Module amont", "Allocation", "Statut allocation",
                "Statut tracabilite", "Meilleur candidat", "Confiance", "Commentaire"],
               arows, widths=[14, 20, 24, 30, 18, 24, 20, 11, 50])

    # ------------------------------------------------------- Liens a auditer
    srows = [[ch["label"]] + list(r.values()) for ch in chains for r in ch["suspicious"]]
    if srows:
        _sheet(wb, "Liens a auditer",
               ["Chaine", "ID amont", "Module amont", "ID aval", "Module aval",
                "Types de lien", "Confiance", "Verdict LLM", "Commentaire"],
               srows, widths=[14, 20, 24, 20, 24, 22, 11, 18, 55])

    # ------------------------------------------------------------ Fiabilite
    calib = context.get("calibration")
    if calib:
        frows = []
        for chain, fit in calib.get("chaines", {}).items():
            if fit.get("mode") == "rejetee":
                frows.append([chain, "calibration refusee", fit["n_exemples"],
                              fit["n_positifs"], fit.get("auc_validation_croisee", ""),
                              "", "", fit.get("motif", "")])
                continue
            for row in fit.get("fiabilite", []):
                frows.append([chain, row["tranche"], fit["n_exemples"],
                              fit["n_positifs"], fit.get("auc_validation_croisee", ""),
                              row["couples"], row["vrais_liens"], row["taux_observe"]])
        if frows:
            _sheet(wb, "Fiabilite",
                   ["Chaine", "Tranche de confiance", "Couples d'apprentissage",
                    "Liens reels", "AUC (validation croisee)", "Couples dans la tranche",
                    "Vrais liens observes", "Taux observe / motif"],
                   frows, widths=[14, 22, 20, 12, 20, 20, 18, 60])

    # ------------------------------------------------------------- Modules
    mod_rows = [[m.level, m.name, m.prefix, len(m.requirements), m.baseline,
                 m.export_time, m.path, m.file] for m in context["modules"]]
    _sheet(wb, "Perimetre",
           ["Niveau", "Module", "Prefixe", "Objets", "Baseline", "Export UTC",
            "Chemin DOORS", "Fichier"],
           mod_rows, widths=[10, 36, 18, 9, 22, 22, 60, 50])

    # ----------------------------------------------------------- Parametres
    prows = [["Date d'execution", context["timestamp"]],
             ["Machine", context["host"]],
             ["Modele embeddings", context["embed_model"]],
             ["Modele juge", context["judge_model"]],
             ["LLM actif", "oui" if context["llm"] else "non"],
             ["Appels LLM", json.dumps(context.get("llm_calls", {}))],
             ["Calibration", context.get("calibration_note", "poids par defaut")],
             ["Seuil proposition", cfg["matching"]["thresholds"]["propose"]],
             ["Seuil a verifier", cfg["matching"]["thresholds"]["weak"]],
             ["Top-k", cfg["matching"]["topk"]],
             ["Pool de candidats", cfg["matching"]["candidate_pool"]],
             ["Poids", json.dumps(cfg["matching"]["weights"])]]
    _sheet(wb, "Parametres", ["Parametre", "Valeur"], prows, widths=[28, 90])

    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)
    return path


def write_markdown(path: str, chains: List[Dict], context: Dict) -> str:
    lines = ["# Audit de tracabilite et de couverture", "",
             f"- Execution : {context['timestamp']}",
             f"- Modeles locaux : `{context['embed_model']}` (embeddings), "
             f"`{context['judge_model']}` (juge)" if context["llm"]
             else "- Mode deterministe (sans LLM)",
             f"- Modules analyses : {len(context['modules'])}", ""]
    for ch in chains:
        s, rs, ev = ch["summary"], ch["reuse"], ch["eval"]
        lines += [f"## Chaine {s['chaine']}", "",
                  "| Indicateur | Valeur |", "|---|---|",
                  f"| Exigences amont | {s['exigences_amont']} |",
                  f"| Deja tracees dans DOORS | {s['deja_tracees']} "
                  f"({s['taux_tracabilite_existant']:.0%}) |",
                  f"| Propositions de chainage | {s['propositions_de_chainage']} |",
                  f"| Couverture apres proposition | "
                  f"{s['taux_couverture_apres_proposition']:.0%} |",
                  f"| Ecarts sans candidat | {s['ecarts_sans_candidat']} |",
                  f"| Couverture partielle detectee | {s['couverture_partielle']} |",
                  f"| Taux de reuse estime | {rs['taux_reuse_estime']:.0%} |",
                  f"| Rappel@5 vs liens DOORS | {ev.get('rappel@5', 'n/a')} |",
                  f"| Precision@1 vs liens DOORS | {ev.get('precision@1', 'n/a')} |",
                  f"| Liens existants a auditer | {len(ch['suspicious'])} |",
                  f"| Allocations non tracees | {len(ch['allocation'])} |", ""]
        gaps = [r for r in ch["results"] if r.status == "ECART_SANS_CANDIDAT"][:10]
        if gaps:
            lines += ["### Exemples d'ecarts", ""]
            for r in gaps:
                lines.append(f"- `{r.src.rid}` — {_cut(r.src.text, 160)}")
            lines.append("")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines))
    return path


def timestamp() -> str:
    return dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
